"""An unresolved tolerance may have fallback numeric limits, but not a pass result."""
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from fai_template_writer import fill_fai_template


@dataclass
class Char:
    char_number: int = 1
    reference_location: str = "P1-R1C1"
    nominal: float = 2.0
    lsl: float | None = 1.98
    usl: float | None = 2.02
    type: str = "LINEAR"
    tooling: str = "CALIPER"
    comments: str = ""
    metadata: dict = field(default_factory=dict)


def _write_template(path: Path, r3: bool) -> None:
    wb = Workbook()
    ws = wb.active
    if r3:
        ws.title = "FAI FORM"
        ws["A3"] = "FIRST ARTICLE INSPECTION (FAI)"
        ws["A22"] = "CHAR"
        ws["C22"] = "REQUIREMENT"
        ws["G22"] = "SUPPLIER INSPECTION RESULT"
        ws["J22"] = "INSPECTION RESULT"
    else:
        headers = ["Char Number", "Reference Location", "Requirement LSL",
                   "Requirement Nominal", "Requirement USL", "Type",
                   "Inspection Actual", "In Spec", "Tooling Used", "Comments"]
        ws.append(headers)
    wb.save(path)


def _output_result(tmp_path, char: Char, r3: bool):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "filled.xlsx"
    _write_template(source, r3)
    fill_fai_template(source, [char], output)
    wb = load_workbook(output, data_only=False)
    cell = wb["FAI FORM"]["K24"] if r3 else wb.active["H2"]
    result = cell.value
    wb.close()
    return result


def test_numeric_defaults_with_review_flag_cannot_pass_generic(tmp_path):
    char = Char(metadata={"tolerance_review_required": True, "tolerance_status": "review_required"})
    assert _output_result(tmp_path, char, r3=False) == "REVIEW REQUIRED"


def test_numeric_defaults_with_review_flag_cannot_pass_r3(tmp_path):
    char = Char(metadata={"tolerance_review_required": True, "tolerance_status": "review_required"})
    assert _output_result(tmp_path, char, r3=True) == "REVIEW REQUIRED"


def test_missing_limit_cannot_pass_even_without_metadata(tmp_path):
    assert _output_result(tmp_path, Char(lsl=None), r3=False) == "REVIEW REQUIRED"


def test_clear_limits_keep_inclusive_formula(tmp_path):
    assert ">=C2" in _output_result(tmp_path, Char(), r3=False)
    assert "<=E2" in _output_result(tmp_path, Char(), r3=False)
    assert ">=C24" in _output_result(tmp_path, Char(), r3=True)
    assert "<=E24" in _output_result(tmp_path, Char(), r3=True)
