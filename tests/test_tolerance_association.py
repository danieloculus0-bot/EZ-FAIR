from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook

import extractor_engine
import extractor_fraction_weld_patch
from fai_template_writer import fill_fai_template
from tolerance_association import associate, nominal_matches


def _drawing(path: Path, text: str):
    doc = fitz.open()
    page = doc.new_page(width=1000, height=250)
    page.insert_text((30, 80), text, fontsize=11)
    doc.save(path)
    doc.close()


@pytest.mark.parametrize("engine", [extractor_engine, extractor_fraction_weld_patch])
def test_same_line_two_nominals_get_their_own_explicit_tolerances(tmp_path, engine):
    drawing = tmp_path / "two.pdf"
    _drawing(drawing, "WIDTH 1.250 +.010/-.005 LENGTH 2.500 +.030/-.010")
    results = engine.extract_pdf_dimensions(drawing)
    first = next(c for c in results if c.nominal == 1.25)
    second = next(c for c in results if c.nominal == 2.5)
    assert (first.lsl, first.usl) == (1.245, 1.26)
    assert (second.lsl, second.usl) == (2.49, 2.53)
    assert first.metadata["tolerance_source"] == "explicit"
    assert second.metadata["tolerance_source"] == "explicit"
    assert not first.metadata["review_required"]
    assert not second.metadata["review_required"]
    assert not any(c.nominal in (0.01, 0.005, 0.03) for c in results)


@pytest.mark.parametrize("engine", [extractor_engine, extractor_fraction_weld_patch])
def test_unattached_tolerance_does_not_get_borrowed(tmp_path, engine):
    drawing = tmp_path / "unattached.pdf"
    _drawing(drawing, "WIDTH 1.250 LENGTH 2.500 +.030/-.010")
    results = engine.extract_pdf_dimensions(drawing)
    first = next(c for c in results if c.nominal == 1.25)
    second = next(c for c in results if c.nominal == 2.5)
    assert first.lsl is None and first.usl is None
    assert first.metadata["review_required"] is True
    assert (second.lsl, second.usl) == (2.49, 2.53)


def test_multiple_adjacent_tolerances_remain_unresolved():
    line = "1.000 +.010/-.020 +.030/-.040 2.000"
    nominals = nominal_matches(line, extractor_engine.DIMENSION_PATTERN,
                               extractor_engine.TOLERANCE_PATTERN)
    assert len(nominals) == 2
    result = associate(line, nominals[0], nominals, extractor_engine.TOLERANCE_PATTERN, 1.0)
    assert result.source == "unresolved"
    assert result.lower is None and result.upper is None


def test_unresolved_limits_are_not_passable_in_excel(tmp_path):
    drawing = tmp_path / "unresolved.pdf"
    _drawing(drawing, "WIDTH 1.250 LENGTH 2.500 +.030/-.010")
    chars = extractor_engine.extract_pdf_dimensions(drawing)
    wb = Workbook()
    ws = wb.active
    ws.append(["Char Number", "Reference Location", "Requirement LSL",
               "Requirement Nominal", "Requirement USL", "Type",
               "company Actual", "In Spec", "Tooling Used", "Comments"])
    template = tmp_path / "template.xlsx"
    wb.save(template)
    output = fill_fai_template(template, chars, tmp_path / "out.xlsx")
    result = load_workbook(output)
    rows = list(result.active.iter_rows(min_row=2, values_only=True))
    first = next(row for row in rows if row[3] == 1.25)
    assert first[2] is None and first[4] is None
    assert first[7] == "REVIEW REQUIRED"
