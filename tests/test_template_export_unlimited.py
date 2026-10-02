from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook

from fai_template_writer import fill_fai_template
from project_store import ProjectMetadata


@dataclass
class Item:
    char_number: int
    reference_location: str
    lsl: float
    nominal: float
    usl: float
    type: str
    page_index: int = 0
    rect: tuple[float, float, float, float] = (0, 0, 1, 1)
    raw_text: str = ""
    tooling: str = "CALIPER"
    comments: str = ""
    actual: str = ""
    metadata: dict = field(default_factory=dict)


def _make_structured_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "FAI FORM"
    ws["A3"] = "FIRST ARTICLE INSPECTION"
    ws["A22"] = "CHAR. NUMBER"
    ws["C22"] = "REQUIREMENT"
    ws["G22"] = "Inspection Result"
    ws["J22"] = "Inspection Result"
    ws["B6"] = ""
    ws["K6"] = ""
    ws["K8"] = ""
    ws["K10"] = ""
    for row in range(24, 49):
        ws.cell(row, 1, row - 23)
        ws.merge_cells(start_row=row, start_column=11, end_row=row, end_column=12)
        ws.cell(row, 11, f'=IF(J{row}="","",IF(AND(J{row}>=C{row},J{row}<=E{row}),"X",""))')
    wb.create_sheet("CHARACTERISTICS")
    wb.create_sheet("TOOLING")
    wb.create_sheet("ATTRIBUTE")
    wb.save(path)


def test_structured_template_export_has_no_characteristic_row_limit(tmp_path: Path) -> None:
    template = tmp_path / "approved_template.xlsx"
    output = tmp_path / "packet.xlsx"
    _make_structured_template(template)
    rows = [
        Item(i, f"P1-R1C1", i - 0.01, float(i), i + 0.01, "LINEAR", comments=f"Row {i}")
        for i in range(1, 61)
    ]

    fill_fai_template(
        template,
        rows,
        output,
        metadata=ProjectMetadata(part_no="P-100", part_name="Bracket", drawing_no="D-100", revision="C"),
    )

    wb = load_workbook(output, data_only=False)
    ws = wb["FAI FORM"]
    assert ws["B6"].value == "P-100"
    assert ws["K6"].value == "Bracket"
    assert ws["K8"].value == "D-100"
    assert ws["K10"].value == "C"
    assert ws["A83"].value == 60
    assert ws["D83"].value == 60.0
    assert ws["C83"].value == 59.99
    assert ws["E83"].value == 60.01
    assert ws["M83"].value == "CALIPER"
    assert ws["N83"].value == "Row 60"
    assert "J83" in ws["K83"].value
    assert any(str(rng) == "K83:L83" for rng in ws.merged_cells.ranges)
    assert str(ws.print_area).endswith("$A$1:$N$83")
    wb.close()
